"""Stress test: every layout and block, both themes, hostile content.

  python tests/stress_test.py [--renderer auto|office|libreoffice|preview] [--keep DIR]

For each generated spec it checks that:
  1. build() does not crash and produces the file,
  2. the file re-opens with python-pptx / python-docx,
  3. page images were produced by the requested renderer,
  4. (if an OOXML schema validator is available, see XSD_VALIDATOR) the XML is schema-valid.
Exit code 0 only if every case passes. Run it on the target machine with
--renderer office to prove PowerPoint and Word open every file without complaint.
"""
from __future__ import annotations

import argparse
import json
import os
import subprocess
import sys
import tempfile
import time
import traceback
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from slidesmith.build import build  # noqa: E402

XSD_VALIDATOR = os.environ.get("XSD_VALIDATOR", "/mnt/skills/public/docx/scripts/office/validate.py")

LONG = ("This sentence is deliberately long so that the renderer has to wrap it across several "
        "lines and possibly shrink it, which is exactly the situation that breaks naive layouts. ")
UNI = "Ünïcödé — κ=0.82 · 日本語テキスト · émojis 🚀✅ · “quotes” & <angle> 'apos'"
URL = "https://example.com/a/very/long/path/without/any/spaces/that/cannot/wrap/naturally/at/all"


def chart(kind, n=6, series=1, neg=False, highlight=None, fmt="0.0"):
    cats = [f"Category {i + 1}" for i in range(n)]
    ser = [{"name": f"Series {s + 1}",
            "values": [round(((i + 1) * (s + 2) % 7 - (3 if neg else 0)) * 0.37, 2) for i in range(n)]}
           for s in range(series)]
    if kind in ("pie", "doughnut"):
        ser = [{"name": "Share", "values": [0.4, 0.25, 0.15, 0.12, 0.08][:n]}]
        cats = cats[:len(ser[0]["values"])]
        fmt = "0%"
    c = {"type": kind, "categories": cats, "series": ser, "number_format": fmt}
    if highlight is not None:
        c["highlight"] = highlight
    return c


def deck_cases(theme):
    slides = [
        {"layout": "title", "kicker": "Stress · " + theme, "title": "Stress test deck " + UNI,
         "subtitle": LONG},
        {"layout": "title", "title": "Minimal title"},
        {"layout": "section", "number": "01", "title": "Section " + UNI, "subtitle": LONG},
        {"layout": "section", "title": ""},
        {"layout": "bullets", "title": "Bullets with emphasis, subs and a takeaway " + UNI,
         "bullets": [{"text": "**Bold lead** then text " + UNI, "sub": ["sub one", "sub two " + URL]},
                     LONG, "short", URL, "*italic* bullet"],
         "takeaway": "Takeaway " + LONG, "notes": "Line 1\nLine 2 " + UNI},
        {"layout": "bullets", "title": "Overflowing bullets must be reported, not crash",
         "bullets": [LONG * 2] * 8},
        {"layout": "two_column", "style": "plain", "title": "Plain columns",
         "left": {"heading": "Left " + UNI, "body": LONG, "bullets": ["a", "b"]},
         "right": {"heading": "Right", "bullets": [LONG, URL]}},
        {"layout": "two_column", "style": "compare", "title": "Compare columns",
         "left": {"heading": "Option A", "bullets": ["x"] * 5}, "right": {"heading": "Option B", "bullets": []}},
        {"layout": "two_column", "style": "before_after", "title": "Before and after",
         "left": {"heading": "Before", "body": LONG * 3}, "right": {"heading": "After", "body": "Short."}},
        {"layout": "big_number", "title": "Big number with context", "number": "$4.2M",
         "label": "Label " + UNI, "context": LONG * 2},
        {"layout": "big_number", "title": "Big number, very long figure", "number": "1,234,567,890",
         "label": "No context"},
        {"layout": "kpis", "title": "Four KPIs with long values",
         "items": [{"value": "$1,234,567", "label": LONG, "delta": "+12%"},
                   {"value": "38%", "label": "Short", "delta": "-4 pts"},
                   {"value": "0.82", "label": UNI, "delta": "-71%", "sentiment": "positive"},
                   {"value": "1.7×", "label": "Ratio", "delta": "flat", "sentiment": "neutral"}],
         "takeaway": LONG},
        {"layout": "kpis", "title": "Two KPIs", "items": [{"value": "1", "label": "a"}, {"value": "2", "label": "b"}]},
    ]
    for kind in ("column", "bar", "line", "pie", "doughnut", "stacked_column", "stacked_bar"):
        slides.append({"layout": "chart", "title": f"{kind} chart",
                       "chart": chart(kind, highlight=1 if kind in ("bar", "column", "pie") else None),
                       "takeaway": LONG if kind in ("bar", "pie") else None, "source": "Synthetic " + UNI})
    slides += [
        {"layout": "chart", "title": "Negative values and many categories",
         "chart": chart("column", n=14, neg=True)},
        {"layout": "chart", "title": "Three-series clustered column, money format",
         "chart": chart("column", n=5, series=3, fmt="$#,##0.0,,\"M\"")},
        {"layout": "chart", "title": "Line with 12 points, no labels",
         "chart": {**chart("line", n=12, series=2), "data_labels": False, "y_title": "Rate"}},
        {"layout": "chart", "title": "All zeros",
         "chart": {"type": "column", "categories": ["a", "b", "c"], "series": [{"name": "z", "values": [0, 0, 0]}]}},
        {"layout": "table", "title": "Max table 7×10 with highlight " + UNI,
         "columns": [f"Col {i}" for i in range(7)],
         "rows": [[f"r{r}c{c}" if c else f"Row {r} {UNI if r == 3 else ''}" for c in range(7)] for r in range(10)],
         "highlight_row": 2, "source": "Synthetic"},
        {"layout": "table", "title": "Table with empty cells and long text",
         "columns": ["Name", "Value", "Comment"], "rows": [["", "1.0", LONG], ["x", "", URL]]},
        {"layout": "timeline", "title": "Six-step timeline",
         "steps": [{"label": f"Q{i}", "title": f"Step {i} " + ("long title " * (i % 3)), "detail": LONG if i % 2 else None}
                   for i in range(1, 7)]},
        {"layout": "timeline", "title": "Three steps", "steps": [{"label": "a", "title": "x"}] * 3},
        {"layout": "quote", "quote": "Quote " + UNI + " " + LONG, "attribution": "Someone, Somewhere"},
        {"layout": "quote", "quote": "Short quote."},
        {"layout": "closing", "title": "Closing " + UNI, "bullets": [LONG, "two", "three", "four", "five"],
         "contact": "team@example.com"},
        {"layout": "closing", "title": "Closing without bullets"},
    ]
    for s in slides:  # drop None values (optional fields)
        for k in [k for k, v in s.items() if v is None]:
            del s[k]
    return {"kind": "deck", "title": "Stress " + UNI, "theme": theme, "author": "QA", "date": "2026",
            "footer": "Footer " + UNI, "slides": slides}


def doc_cases(theme):
    blocks = [
        {"type": "callout", "kind": "key", "title": "Key " + UNI, "text": "**Bold** and *italic* " + LONG},
        {"type": "kpis", "items": [{"value": "$1.2M", "label": "a", "delta": "+1"},
                                   {"value": "38%", "label": UNI, "delta": "-2"},
                                   {"value": "0.82", "label": LONG, "sentiment": "neutral", "delta": "0"},
                                   {"value": "7", "label": "d"}]},
        {"type": "heading", "level": 1, "text": "Heading one " + UNI},
        {"type": "paragraph", "text": (LONG * 12) + URL},
        {"type": "heading", "level": 2, "text": "Heading two"},
        {"type": "bullets", "items": [{"text": "**A** " + LONG, "sub": ["s1", "s2 " + UNI]}, "b", URL]},
        {"type": "bullets", "ordered": True, "items": ["one", "two", "three"]},
        {"type": "paragraph", "text": "Between lists."},
        {"type": "bullets", "ordered": True, "items": ["restart one", "restart two"]},
        {"type": "heading", "level": 3, "text": "Heading three"},
        {"type": "callout", "kind": "info", "text": "Info without title."},
        {"type": "callout", "kind": "warning", "title": "Warning", "text": LONG * 2},
        {"type": "table", "columns": [f"Column {i}" for i in range(7)],
         "rows": [[f"{r * 1.5:.1f}" if c else f"Row {r} {UNI if r == 5 else ''}" for c in range(7)] for r in range(30)],
         "caption": "Table 1. Thirty rows to force a page split " + UNI},
        {"type": "page_break"},
    ]
    for kind in ("column", "bar", "line", "pie", "doughnut", "stacked_column", "stacked_bar"):
        blocks += [{"type": "paragraph", "text": f"Lead-in for the {kind} chart."},
                   {"type": "chart", "chart": chart(kind, neg=kind == "column", highlight=0 if kind == "bar" else None),
                    "caption": f"Figure: {kind}"}]
    blocks += [
        {"type": "chart", "chart": chart("column", n=5, series=3, fmt="$#,##0.0,,\"M\"")},
        {"type": "quote", "text": "Quote " + UNI, "attribution": "Someone"},
        {"type": "quote", "text": "Quote without attribution."},
        {"type": "heading", "level": 1, "text": "Last section"},
        {"type": "paragraph", "text": ""},
        {"type": "paragraph", "text": "End."},
    ]
    return {"kind": "doc", "title": "Stress document " + UNI, "subtitle": LONG, "theme": theme,
            "author": "QA", "date": "2026", "header": "Header " + UNI, "blocks": blocks}


def reopen(path: Path):
    if path.suffix == ".pptx":
        from pptx import Presentation
        return len(Presentation(str(path)).slides)
    from docx import Document
    return len(Document(str(path)).paragraphs)


def xsd_validate(path: Path) -> str:
    if not Path(XSD_VALIDATOR).exists():
        return "skipped (no validator)"
    r = subprocess.run([sys.executable, XSD_VALIDATOR, str(path)], capture_output=True, text=True,
                       cwd=str(Path(XSD_VALIDATOR).parent.parent))
    out = (r.stdout + r.stderr).strip()
    return "valid" if "PASSED" in out else "INVALID: " + out[-600:]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--renderer", default="auto")
    ap.add_argument("--keep", help="keep outputs in this folder")
    a = ap.parse_args()
    work = Path(a.keep) if a.keep else Path(tempfile.mkdtemp(prefix="slidesmith-stress-"))
    work.mkdir(parents=True, exist_ok=True)
    cases = []
    for theme in ("default", "warm"):
        cases.append((f"deck-{theme}", deck_cases(theme)))
        cases.append((f"doc-{theme}", doc_cases(theme)))
    failures = 0
    for name, spec in cases:
        sp = work / f"{name}.{'deck' if spec['kind'] == 'deck' else 'doc'}.json"
        sp.write_text(json.dumps(spec, ensure_ascii=False, indent=1), encoding="utf-8")
        t0 = time.time()
        problems = []
        try:
            res = build(sp, work / name, renderer=a.renderer)
            if not res.output or not Path(res.output).exists():
                problems.append("no output file: " + json.dumps(res.summary)[:300])
            else:
                n = reopen(Path(res.output))
                if not res.pages:
                    problems.append("no page images")
                if a.renderer != "auto" and res.renderer != a.renderer:
                    problems.append(f"renderer was {res.renderer}, wanted {a.renderer}")
                x = xsd_validate(Path(res.output))
                if x.startswith("INVALID"):
                    problems.append(x)
            errs = [i for i in res.summary.get("issues", []) if i["severity"] == "error"]
            detail = (f"{len(res.pages)} pages via {res.renderer}, score {res.summary['score']}, "
                      f"{len(errs)} design errors reported, xml {x if res.output else '-'}")
        except Exception:  # noqa: BLE001
            problems.append("CRASH:\n" + traceback.format_exc())
            detail = ""
        status = "PASS" if not problems else "FAIL"
        failures += bool(problems)
        print(f"[{status}] {name:14s} {time.time() - t0:5.1f}s  {detail}")
        for p in problems:
            print("        " + p.replace("\n", "\n        "))
    print(f"\n{len(cases) - failures}/{len(cases)} passed. Outputs in {work}")
    sys.exit(1 if failures else 0)


if __name__ == "__main__":
    main()
