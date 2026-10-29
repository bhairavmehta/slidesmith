"""Checks that the figure-provenance guard flags wrong numbers and accepts legitimate ones.

  python tests/test_provenance.py
"""
from __future__ import annotations

import json
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from slidesmith import provenance  # noqa: E402
from slidesmith.schema import load_spec  # noqa: E402

FAILS = 0


def expect(name, cond, detail=""):
    global FAILS
    print(f"[{'PASS' if cond else 'FAIL'}] {name}" + (f"  -  {detail}" if detail and not cond else ""))
    FAILS += not cond


def flagged(spec_dict, files: dict[str, str | bytes]) -> list[str]:
    with tempfile.TemporaryDirectory() as d:
        d = Path(d)
        for name, content in files.items():
            (d / name).write_bytes(content if isinstance(content, bytes) else content.encode("utf-8"))
        sp = d / "x.deck.json"
        sp.write_text(json.dumps(spec_dict), encoding="utf-8")
        issues = provenance.check(load_spec(spec_dict), sp)
        return [i["message"] for i in issues if i["rule"] in ("unverified_figure", "source_missing")]


def deck(text, chart_values=None, sources=("src.md",)):
    slides = [{"layout": "bullets", "title": "Results this quarter improved",
               "bullets": [text]}]
    if chart_values:
        slides.append({"layout": "chart", "title": "Chart of the values over time",
                       "chart": {"type": "column", "categories": [str(i) for i in range(len(chart_values))],
                                 "series": [{"name": "v", "values": chart_values}]}})
    return {"kind": "deck", "title": "t", "sources": list(sources), "slides": slides}


SRC = "Revenue $4.2M. Margin 38.4%. Kappa 0.8213. Customers 12,450. Growth 0.17\n"

expect("exact figures pass", not flagged(deck("Revenue **$4.2M**, customers 12,450"), {"src.md": SRC}))
expect("percent written as 38% from 38.4% passes (rounded)", not flagged(deck("Margin 38%"), {"src.md": SRC}))
expect("0.82 from 0.8213 passes (rounded)", not flagged(deck("Kappa 0.82"), {"src.md": SRC}))
expect("fraction 0.17 shown as 17% passes", not flagged(deck("Growth 17%"), {"src.md": SRC}))
expect("$4,200,000 shown for $4.2M passes", not flagged(deck("Revenue $4,200,000"), {"src.md": SRC}))
expect("wrong number is flagged", any("0.85" in m for m in flagged(deck("Kappa 0.85"), {"src.md": SRC})))
expect("wrong money is flagged", any("4.7" in m for m in flagged(deck("Revenue $4.7M"), {"src.md": SRC})))
expect("transposed digits flagged", any("12540" in m for m in flagged(deck("Customers 12,540"), {"src.md": SRC})))
expect("wrong chart value flagged", any("chart" in m for m in flagged(deck("ok", [4.2, 9.9]), {"src.md": SRC})))
expect("years and small counts ignored", not flagged(deck("In 2026 we ran 3 pilots"), {"src.md": SRC}))
expect("missing source reported", any("not found" in m for m in flagged(deck("x", sources=("nope.xlsx",)), {})))

# binary source formats
with tempfile.TemporaryDirectory() as d:
    d = Path(d)
    from docx import Document
    doc = Document(); doc.add_paragraph("Churn fell to 3.7% in Q3."); doc.save(d / "s.docx")
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation(); s = prs.slides.add_slide(prs.slide_layouts[6])
    s.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text_frame.text = "NPS 47.5"
    prs.save(d / "s.pptx")
    import openpyxl
    wb = openpyxl.Workbook(); wb.active.append(["Net revenue", 1234567.89]); wb.save(d / "s.xlsx")
    from PIL import Image
    pdf_ok = True
    try:
        from reportlab.pdfgen import canvas
        c = canvas.Canvas(str(d / "s.pdf")); c.drawString(72, 720, "Headcount 1,830"); c.save()
    except ImportError:
        pdf_ok = False
    files = {n: (d / n).read_bytes() for n in ("s.docx", "s.pptx", "s.xlsx")}
    srcs = ["s.docx", "s.pptx", "s.xlsx"]
    text = "Churn 3.7%, NPS 47.5, net revenue $1.23M"
    if pdf_ok:
        files["s.pdf"] = (d / "s.pdf").read_bytes(); srcs.append("s.pdf"); text += ", headcount 1,830"
    got = flagged(deck(text, sources=srcs), files)
    expect(f"docx/pptx/xlsx{'/pdf' if pdf_ok else ''} sources read", not got, "; ".join(got))

print(f"\n{'All provenance checks passed.' if not FAILS else f'{FAILS} check(s) FAILED.'}")
sys.exit(1 if FAILS else 0)
