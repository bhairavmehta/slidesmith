"""Document spec -> .docx with a consistent typographic system.

Uses real Word styles (Title/Heading 1-3/List Bullet...) so the navigation pane,
table of contents and restyling in Word all keep working.
"""
from __future__ import annotations

import copy
import re
import tempfile
from pathlib import Path

from docx import Document as WordDocument
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from . import markup
from .charts_mpl import render_chart_png
from .schema import (
    BulletsBlock, CalloutBlock, ChartBlock, Document, HeadingBlock, KPIBlock, PageBreakBlock,
    ParagraphBlock, QuoteBlock, TableBlock, bullet_subs, bullet_text,
)
from .theme import Theme, mix

NUMERIC = re.compile(r"^[\s$€£¥+\-−~≈<>]*[\d.,]+\s*(%|x|pp|bps|[KMBkmb]|ms|s|pts?)?\s*$")


def _rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h.lstrip("#").upper())


# --------------------------------------------------------------------------- #
# low-level XML helpers
# --------------------------------------------------------------------------- #
# Word rejects ("unreadable content") children that are out of schema order, so every
# element we add is inserted at its schema position, never simply appended.
PPR_ORDER = ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr",
             "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap",
             "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd",
             "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc",
             "textDirection", "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr",
             "sectPr", "pPrChange"]
TBLPR_ORDER = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize",
               "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd",
               "tblLayout", "tblCellMar", "tblLook", "tblCaption", "tblDescription", "tblPrChange"]
TCPR_ORDER = ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar",
              "textDirection", "tcFitText", "vAlign", "hideMark", "headers", "cellIns", "cellDel",
              "cellMerge", "tcPrChange"]
BORDER_ORDER = {"pBdr": ["top", "left", "bottom", "right", "between", "bar"],
                "tcBorders": ["top", "left", "start", "bottom", "right", "end", "insideH", "insideV",
                              "tl2br", "tr2bl"],
                "tblBorders": ["top", "left", "start", "bottom", "right", "end", "insideH", "insideV"],
                "tblCellMar": ["top", "left", "start", "bottom", "right", "end"]}


def _local(el) -> str:
    return el.tag.split("}")[-1]


def set_child(parent, child, order: list[str]):
    """Insert child at its schema position, replacing an existing element of the same name."""
    name = _local(child)
    for old in parent.findall(qn(f"w:{name}")):
        parent.remove(old)
    rank = order.index(name)
    for i, sib in enumerate(list(parent)):
        n = _local(sib)
        if n in order and order.index(n) > rank:
            parent.insert(i, child)
            return child
    parent.append(child)
    return child


def _border_el(edge: str, color: str | None, size_eighths: int = 0, space: int = 4):
    b = OxmlElement(f"w:{edge}")
    if color is None:
        b.set(qn("w:val"), "nil")
    else:
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), str(size_eighths))
        b.set(qn("w:space"), str(space))
        b.set(qn("w:color"), color)
    return b


def para_border(p, edge: str, color: str, size_eighths: int = 8, space: int = 4):
    pPr = p._p.get_or_add_pPr()
    pbdr = pPr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = set_child(pPr, OxmlElement("w:pBdr"), PPR_ORDER)
    set_child(pbdr, _border_el(edge, color, size_eighths, space), BORDER_ORDER["pBdr"])


def cell_borders(cell, **edges):
    """edges: top/left/bottom/right = (color, size_eighths); omitted edges are cleared."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcb = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        spec = edges.get(edge)
        tcb.append(_border_el(edge, spec[0], spec[1], 0) if spec else _border_el(edge, None))
    set_child(tcPr, tcb, TCPR_ORDER)


def cell_shade(cell, color: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    set_child(tcPr, shd, TCPR_ORDER)


def cell_margins(table, top=80, bottom=80, left=120, right=120):
    mar = OxmlElement("w:tblCellMar")
    for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{k}")
        e.set(qn("w:w"), str(v))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    set_child(table._tbl.tblPr, mar, TBLPR_ORDER)


def keep_row_together(table):
    """Word moves an unsplittable row to the next page instead of breaking the box."""
    for row in table.rows:
        trPr = row._tr.get_or_add_trPr()
        if trPr.find(qn("w:cantSplit")) is None:
            trPr.append(OxmlElement("w:cantSplit"))


def no_table_borders(table):
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b.append(_border_el(edge, None))
    set_child(table._tbl.tblPr, b, TBLPR_ORDER)


def add_field(run, instr: str):
    for tag, text in (("begin", None), (None, instr), ("separate", None), (None, "1"), ("end", None)):
        if tag:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), tag)
            run._r.append(fc)
        elif text == instr:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = f" {instr} "
            run._r.append(it)
        else:
            t = OxmlElement("w:t")
            t.text = text
            run._r.append(t)


def keep_with_next(p):
    p.paragraph_format.keep_with_next = True


# --------------------------------------------------------------------------- #
class DocRenderer:
    def __init__(self, doc: Document, theme: Theme):
        self.spec = doc
        self.t = theme
        self.sz = theme.doc["size"]
        self.issues: list[dict] = []
        self.d = WordDocument()
        self._tmp = tempfile.TemporaryDirectory()

    def issue(self, idx: int, severity: str, rule: str, message: str, hint: str = ""):
        self.issues.append({"where": f"block {idx}", "index": idx, "severity": severity,
                            "rule": rule, "message": message, "hint": hint})

    # ---------------- setup ----------------
    def _styles(self):
        t, sz = self.t, self.sz
        st = self.d.styles
        normal = st["Normal"]
        normal.font.name = t.body_font
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), t.body_font)
        normal.font.size = Pt(sz["body"])
        normal.font.color.rgb = _rgb(t.c("ink"))
        pf = normal.paragraph_format
        pf.space_after = Pt(t.doc.get("para_after_pt", 8))
        pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        pf.line_spacing = t.doc.get("line_spacing", 1.2)

        specs = {
            "Title": (sz["title"], t.c("ink"), 0, 4),
            "Subtitle": (sz["subtitle"], t.c("muted"), 0, 6),
            "Heading 1": (sz["h1"], t.c("primary"), 22, 8),
            "Heading 2": (sz["h2"], t.c("ink"), 16, 6),
            "Heading 3": (sz["h3"], t.c("accent"), 12, 4),
        }
        for name, (size, color, before, after) in specs.items():
            s = st[name]
            s.font.name = t.heading_font
            rpr = s.element.get_or_add_rPr()
            rf = rpr.find(qn("w:rFonts"))
            if rf is None:  # rFonts must come right after rStyle
                rf = OxmlElement("w:rFonts")
                rs = rpr.find(qn("w:rStyle"))
                rpr.insert(0 if rs is None else list(rpr).index(rs) + 1, rf)
            for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                rf.set(qn(a), t.heading_font)
            for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                if rf.get(qn(a)) is not None:
                    del rf.attrib[qn(a)]
            for sp in rpr.findall(qn("w:spacing")):  # Subtitle ships with letter-spacing
                rpr.remove(sp)
            s.font.size = Pt(size)
            s.font.color.rgb = _rgb(color)
            s.font.bold = name != "Subtitle"
            s.font.italic = False
            s.paragraph_format.space_before = Pt(before)
            s.paragraph_format.space_after = Pt(after)
            s.paragraph_format.line_spacing = 1.05
            if name == "Title":
                # Title style ships with a bottom border; drop it, we draw our own rule.
                ppr = s.element.get_or_add_pPr()
                for b in ppr.findall(qn("w:pBdr")):
                    ppr.remove(b)
        for name in ("List Bullet", "List Bullet 2", "List Number"):
            s = st[name]
            s.font.name = t.body_font
            s.paragraph_format.space_after = Pt(4)

    def _fix_settings(self):
        settings = self.d.settings.element
        zoom = settings.find(qn("w:zoom"))
        if zoom is not None and zoom.get(qn("w:percent")) is None:
            zoom.set(qn("w:percent"), "100")

    def _page(self):
        sec = self.d.sections[0]
        m = Inches(self.t.doc.get("margin_in", 1.0))
        sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = m
        if self.t.doc.get("page", "letter") == "a4":
            sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
        else:
            sec.page_width, sec.page_height = Inches(8.5), Inches(11)
        sec.different_first_page_header_footer = True
        muted = self.t.c("muted")
        # running header (not on page 1)
        hp = sec.header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = hp.add_run(self.spec.header or self.spec.title)
        r.font.size = Pt(self.sz["small"])
        r.font.color.rgb = _rgb(muted)
        # page number footer on every page
        for footer in (sec.footer, sec.first_page_footer):
            fp = footer.paragraphs[0]
            fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            r = fp.add_run()
            r.font.size = Pt(self.sz["small"])
            r.font.color.rgb = _rgb(muted)
            add_field(r, "PAGE")

    @property
    def text_width_in(self) -> float:
        sec = self.d.sections[0]
        return (sec.page_width - sec.left_margin - sec.right_margin) / 914400

    # ---------------- inline text ----------------
    def runs(self, p, text: str, size=None, color=None, bold=None, italic=None, font=None):
        for seg, b, it in markup.runs(text):
            r = p.add_run(seg)
            if size:
                r.font.size = Pt(size)
            if color:
                r.font.color.rgb = _rgb(color)
            if font:
                r.font.name = font
            r.font.bold = True if (b or bold) else None
            r.font.italic = True if (it or italic) else None
        return p

    # ---------------- render ----------------
    def render(self, path: str) -> list[dict]:
        self._styles()
        self._fix_settings()
        self._page()
        d, t = self.d, self.t
        cp = d.core_properties
        cp.title = self.spec.title
        if self.spec.author:
            cp.author = self.spec.author

        # title block
        body = d.paragraphs[0] if d.paragraphs else d.add_paragraph()
        body.style = d.styles["Title"]
        body.text = ""
        self.runs(body, self.spec.title)
        if self.spec.subtitle:
            self.runs(d.add_paragraph(style="Subtitle"), self.spec.subtitle)
        meta = " · ".join(v for v in [self.spec.author, self.spec.date] if v)
        rule = d.add_paragraph()
        if meta:
            self.runs(rule, meta, size=self.sz["small"] + 1, color=t.c("muted"))
        rule.paragraph_format.space_after = Pt(18)
        para_border(rule, "bottom", t.c("accent"), size_eighths=12, space=6)

        self._numbering_ids = 0
        for i, b in enumerate(self.spec.blocks, start=1):
            getattr(self, f"_{b.type}")(b, i)
        d.save(path)
        self._tmp.cleanup()
        return self.issues

    def _heading(self, b: HeadingBlock, i: int):
        p = self.d.add_paragraph(style=f"Heading {b.level}")
        self.runs(p, b.text)
        keep_with_next(p)

    def _paragraph(self, b: ParagraphBlock, i: int):
        self.runs(self.d.add_paragraph(), b.text)

    def _bullets(self, b: BulletsBlock, i: int):
        num_id = self._new_numbering() if b.ordered else None
        for item in b.items:
            p = self.d.add_paragraph(style="List Number" if b.ordered else "List Bullet")
            if num_id:
                self._set_num(p, num_id, 0)
            self.runs(p, bullet_text(item))
            for s in bullet_subs(item):
                sp = self.d.add_paragraph(style="List Bullet 2")
                self.runs(sp, s, color=self.t.c("muted"))
        last = self.d.paragraphs[-1]
        last.paragraph_format.space_after = Pt(self.t.doc.get("para_after_pt", 8) + 2)

    def _new_numbering(self) -> int:
        """A fresh w:num that restarts at 1 but reuses List Number's abstract definition."""
        numbering = self.d.part.numbering_part.numbering_definitions._numbering
        style = self.d.styles["List Number"].element
        base_id = int(style.pPr.numPr.numId.val)
        base = [n for n in numbering.findall(qn("w:num")) if int(n.get(qn("w:numId"))) == base_id][0]
        abstract = base.find(qn("w:abstractNumId")).get(qn("w:val"))
        new_id = max(int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))) + 1
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(new_id))
        a = OxmlElement("w:abstractNumId")
        a.set(qn("w:val"), abstract)
        num.append(a)
        ov = OxmlElement("w:lvlOverride")
        ov.set(qn("w:ilvl"), "0")
        so = OxmlElement("w:startOverride")
        so.set(qn("w:val"), "1")
        ov.append(so)
        num.append(ov)
        nums = numbering.findall(qn("w:num"))  # w:num must follow the last w:num
        numbering.insert(list(numbering).index(nums[-1]) + 1, num)
        return new_id

    @staticmethod
    def _set_num(p, num_id: int, level: int):
        pPr = p._p.get_or_add_pPr()
        numPr = OxmlElement("w:numPr")
        il = OxmlElement("w:ilvl")
        il.set(qn("w:val"), str(level))
        ni = OxmlElement("w:numId")
        ni.set(qn("w:val"), str(num_id))
        numPr.append(il)
        numPr.append(ni)
        set_child(pPr, numPr, PPR_ORDER)

    def _callout(self, b: CalloutBlock, i: int):
        t = self.t
        stripe = {"key": t.c("accent"), "info": t.c("primary"), "warning": t.c("negative")}[b.kind]
        tbl = self.d.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        no_table_borders(tbl)
        cell_margins(tbl, 140, 140, 220, 220)
        keep_row_together(tbl)
        cell = tbl.cell(0, 0)
        cell.width = Inches(self.text_width_in)
        cell_shade(cell, mix(stripe, t.c("bg"), 0.9))
        cell_borders(cell, left=(stripe, 24))
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        if b.title:
            self.runs(p, b.title, bold=True, color=t.c("ink"))
            p = cell.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
        else:
            p.paragraph_format.space_after = Pt(0)
        self.runs(p, b.text)
        self._spacer(12)

    def _spacer(self, pts: int = 6):
        sp = self.d.add_paragraph()
        sp.paragraph_format.space_after = Pt(0)
        sp.paragraph_format.line_spacing = Pt(pts)

    def _table(self, b: TableBlock, i: int):
        t = self.t
        if len(b.columns) > 6:
            self.issue(i, "warning", "table_width",
                       f"table has {len(b.columns)} columns; portrait pages fit ~6 comfortably.",
                       "Drop or merge columns, or move detail to an appendix.")
        tbl = self.d.add_table(rows=len(b.rows) + 1, cols=len(b.columns))
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        no_table_borders(tbl)
        cell_margins(tbl, 70, 70, 110, 110)
        lens = [max([len(b.columns[c])] + [len(r[c]) for r in b.rows]) for c in range(len(b.columns))]
        weights = [max(6, min(n, 40)) ** 0.75 for n in lens]
        widths = [self.text_width_in * w / sum(weights) for w in weights]
        numeric = [sum(bool(NUMERIC.match(r[c])) for r in b.rows) >= max(1, len(b.rows) * 0.6)
                   for c in range(len(b.columns))]
        for ri, cells in enumerate([b.columns] + b.rows):
            row = tbl.rows[ri]
            if ri == 0:
                trPr = row._tr.get_or_add_trPr()
                th = OxmlElement("w:tblHeader")
                trPr.append(th)
            for c, txt in enumerate(cells):
                cell = row.cells[c]
                cell.width = Inches(widths[c])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                if ri == 0:
                    cell_borders(cell, bottom=(t.c("primary"), 12))
                else:
                    cell_borders(cell, bottom=(t.c("line"), 4))
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.1
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if numeric[c] and c > 0 else WD_ALIGN_PARAGRAPH.LEFT
                self.runs(p, txt, size=self.sz["body"] - (1 if ri == 0 else 0.5),
                          bold=ri == 0 or None, color=t.c("muted") if ri == 0 else None)
        self._caption(b.caption)

    def _caption(self, text):
        if text:
            p = self.d.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            self.runs(p, text, size=self.sz["small"], color=self.t.c("muted"), italic=True)
        else:
            self._spacer(10)

    def _chart(self, b: ChartBlock, i: int):
        png = str(Path(self._tmp.name) / f"chart_{i}.png")
        render_chart_png(b.chart, self.t, png, width_in=self.text_width_in, height_in=3.2)
        if self.d.paragraphs:  # keep the lead-in text on the same page as the figure
            keep_with_next(self.d.paragraphs[-1])
        p = self.d.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        keep_with_next(p)
        p.add_run().add_picture(png, width=Inches(self.text_width_in))
        self._caption(b.caption)

    def _kpis(self, b: KPIBlock, i: int):
        t = self.t
        n = len(b.items)
        tbl = self.d.add_table(rows=1, cols=n)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        no_table_borders(tbl)
        cell_margins(tbl, 120, 120, 160, 160)
        keep_row_together(tbl)
        for j, k in enumerate(b.items):
            cell = tbl.cell(0, j)
            cell.width = Inches(self.text_width_in / n)
            cell_shade(cell, t.c("surface"))
            cell_borders(cell, top=(t.c("accent") if j == 0 else t.c("primary"), 18),
                         left=(t.c("bg"), 24) if j else None, right=(t.c("bg"), 24) if j < n - 1 else None)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            self.runs(p, k.value, size=self.sz["kpi_value"], bold=True, font=t.heading_font)
            p = cell.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.1
            self.runs(p, k.label, size=self.sz["body"] - 1.5, color=t.c("muted"))
            if k.delta:
                sent = k.sentiment or ("negative" if k.delta.strip().startswith(("-", "−", "▼", "↓"))
                                       else "positive")
                col = {"positive": t.c("positive"), "negative": t.c("negative"), "neutral": t.c("muted")}[sent]
                p = cell.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                self.runs(p, k.delta, size=self.sz["body"] - 2, bold=True, color=col)
        self._spacer(14)

    def _quote(self, b: QuoteBlock, i: int):
        t = self.t
        p = self.d.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2 if b.attribution else 12)
        para_border(p, "left", t.c("accent"), size_eighths=24, space=12)
        self.runs(p, b.text, size=self.sz["body"] + 3, italic=True, font=t.heading_font)
        if b.attribution:
            a = self.d.add_paragraph()
            a.paragraph_format.left_indent = Inches(0.25)
            a.paragraph_format.space_after = Pt(12)
            para_border(a, "left", t.c("accent"), size_eighths=24, space=12)
            self.runs(a, f"— {b.attribution}", size=self.sz["small"] + 1, color=t.c("muted"), bold=True)

    def _page_break(self, b: PageBreakBlock, i: int):
        self.d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def render_doc(doc: Document, theme: Theme, path: str) -> list[dict]:
    return DocRenderer(doc, theme).render(path)
