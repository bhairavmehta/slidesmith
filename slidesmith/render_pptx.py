"""Deck spec -> .pptx.

Every slide is drawn on a blank layout with explicit geometry from the theme.
Text is sized with real font metrics (textfit) and anything that cannot fit at
the theme's minimum size is reported as an issue rather than silently shrunk.
Charts are native PowerPoint charts, so they stay editable.
"""
from __future__ import annotations

import re
from dataclasses import dataclass
from typing import Optional

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import (
    XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_TICK_LABEL_POSITION, XL_TICK_MARK,
)
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from . import markup, textfit
from .schema import (
    BigNumberSlide, BulletsSlide, ChartSlide, ChartSpec, ClosingSlide, Column, Deck, KPISlide,
    QuoteSlide, SectionSlide, TableSlide, TimelineSlide, TitleSlide, TwoColumnSlide,
    bullet_subs, bullet_text,
)
from .theme import Theme, mix

ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
ANCHOR = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}
NUMERIC = re.compile(r"^[\s$€£¥+\-−~≈<>]*[\d.,]+\s*(%|x|pp|bps|[KMBkmb]|ms|s|pts?)?\s*$")


@dataclass
class TP:
    """A paragraph to draw."""
    text: str
    scale: float = 1.0
    bold: bool = False
    italic: bool = False
    color: Optional[str] = None
    bullet: Optional[str] = None      # bullet character, or "#" for auto-numbering
    level: int = 0
    space_before: float = 0.0         # multiples of the base size


def _rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h.lstrip("#").upper())


class SlideCanvas:
    def __init__(self, deck_renderer: "DeckRenderer", slide, index: int):
        self.r = deck_renderer
        self.t: Theme = deck_renderer.theme
        self.s = slide
        self.i = index
        self.sz = self.t.slide["size"]
        self.ls = self.t.slide.get("line_spacing", 1.1)

    # ---------------- primitives ----------------
    def bg(self, color: str):
        fill = self.s.background.fill
        fill.solid()
        fill.fore_color.rgb = _rgb(color)

    def rect(self, x, y, w, h, fill: str, shape=MSO_SHAPE.RECTANGLE, line: Optional[str] = None):
        shp = self.s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        shp.fill.solid()
        shp.fill.fore_color.rgb = _rgb(fill)
        if line:
            shp.line.color.rgb = _rgb(line)
            shp.line.width = Pt(1)
        else:
            shp.line.fill.background()
        # Drop the theme style reference so no renderer adds a default shadow/outline.
        style = shp._element.find(qn("p:style"))
        if style is not None:
            shp._element.remove(style)
        if shp.has_text_frame:
            shp.text_frame.text = ""
        return shp

    def text(self, x, y, w, h, paras: list[TP], *, size: float, min_size: float, color: str,
             font: Optional[str] = None, bold: bool = False, align: str = "l", anchor: str = "t",
             ls: Optional[float] = None, element: str = "text", max_lines: Optional[int] = None,
             report: bool = True) -> textfit.Fit:
        font = font or self.t.body_font
        ls = ls or self.ls
        tf_paras = [
            textfit.Para(markup.plain(p.text), p.scale, self._indent(p), p.bold or bold, p.space_before)
            for p in paras
        ]
        f = textfit.fit(tf_paras, font, w, h, size, min_size, ls, max_lines)
        if not f.fits and report:
            words = sum(len(markup.plain(p.text).split()) for p in paras)
            self.r.issue(self.i, "error", "overflow",
                         f"{element} does not fit its box even at {f.pt:.0f}pt "
                         f"({f.lines} lines, {words} words).",
                         "Shorten the text (cut ~30% of the words) or move detail to speaker notes.")

        box = self.s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = box.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = ANCHOR[anchor]
        for k, p in enumerate(paras):
            para = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            para.alignment = ALIGN[align]
            para.line_spacing = ls
            if p.space_before:
                para.space_before = Pt(p.space_before * f.pt)
            if p.bullet:
                self._bullet(para, p, p.color or color)
            for seg, b, it in markup.runs(p.text):
                run = para.add_run()
                run.text = seg
                fnt = run.font
                fnt.name = font
                fnt.size = Pt(round(f.pt * p.scale, 1))
                fnt.bold = bold or p.bold or b
                fnt.italic = p.italic or it
                fnt.color.rgb = _rgb(p.color or color)
        return f

    def _indent(self, p: TP) -> float:
        if not p.bullet:
            return 0.0
        return 0.34 + 0.34 * p.level

    def _bullet(self, para, p: TP, color: str):
        pPr = para._p.get_or_add_pPr()
        mar = self._indent(p)
        hang = 0.30 if p.level == 0 else 0.26
        pPr.set("marL", str(int(Inches(mar))))
        pPr.set("indent", str(-int(Inches(hang))))
        clr = etree.SubElement(pPr, qn("a:buClr"))
        etree.SubElement(clr, qn("a:srgbClr"), val=(self.t.c("accent") if p.level == 0 else self.t.c("muted")))
        etree.SubElement(pPr, qn("a:buSzPct"), val="100000")
        if p.bullet == "#":
            etree.SubElement(pPr, qn("a:buFont"), typeface=self.t.body_font)
            etree.SubElement(pPr, qn("a:buAutoNum"), type="arabicPeriod")
        else:
            etree.SubElement(pPr, qn("a:buFont"), typeface="Arial")
            etree.SubElement(pPr, qn("a:buChar"), char=p.bullet)

    # ---------------- shared chrome ----------------
    @property
    def W(self): return self.t.slide["width_in"]

    @property
    def H(self): return self.t.slide["height_in"]

    @property
    def mx(self): return self.t.slide["margin_x"]

    @property
    def cw(self): return self.W - 2 * self.mx

    @property
    def ctop(self): return self.t.slide["content_top"]

    @property
    def cbottom(self): return self.H - self.t.slide["margin_bottom"] - 0.25

    def chrome(self, title: str):
        """Background, accent tick, action title, footer, page number."""
        t = self.t
        self.bg(t.c("bg"))
        top = t.slide["margin_top"]
        self.rect(self.mx, top - 0.16, 0.42, 0.055, t.c("accent"))
        self.text(self.mx, top, self.cw, t.slide["title_h"], [TP(title)],
                  size=self.sz["title"], min_size=max(18, self.sz["title"] - 8),
                  color=t.c("ink"), font=t.heading_font, bold=True, ls=1.0,
                  element="title", max_lines=2)
        fy = self.H - 0.45
        footer = self.r.deck.footer or self.r.deck.title
        self.text(self.mx, fy, self.cw - 1, 0.3, [TP(footer)], size=self.sz["small"],
                  min_size=8, color=t.c("muted"), report=False)
        self.text(self.W - self.mx - 1, fy, 1, 0.3, [TP(str(self.i))], size=self.sz["small"],
                  min_size=8, color=t.c("muted"), align="r", report=False)

    def takeaway_bar(self, text: str, y: float, h: float = 0.8):
        t = self.t
        self.rect(self.mx, y, self.cw, h, t.c("surface"))
        self.rect(self.mx, y, 0.08, h, t.c("accent"))
        self.text(self.mx + 0.35, y + 0.1, self.cw - 0.6, h - 0.2, [TP(text)],
                  size=self.sz["takeaway"] + 2, min_size=12, color=t.c("ink"), bold=True,
                  anchor="m", element="takeaway", max_lines=2)

    def bullet_paras(self, items, color: str) -> list[TP]:
        out = []
        for k, b in enumerate(items):
            out.append(TP(bullet_text(b), bullet="•", space_before=0 if k == 0 else 0.55))
            for s in bullet_subs(b):
                out.append(TP(s, scale=0.82, bullet="–", level=1, color=self.t.c("muted"),
                              space_before=0.2))
        return out


class DeckRenderer:
    def __init__(self, deck: Deck, theme: Theme):
        self.deck = deck
        self.theme = theme
        self.issues: list[dict] = []
        self.prs = Presentation()
        self.prs.slide_width = Inches(theme.slide["width_in"])
        self.prs.slide_height = Inches(theme.slide["height_in"])
        self._blank = self.prs.slide_layouts[6]

    def issue(self, slide: int, severity: str, rule: str, message: str, hint: str = ""):
        self.issues.append({"where": f"slide {slide}", "index": slide, "severity": severity,
                            "rule": rule, "message": message, "hint": hint})

    def render(self, path: str) -> list[dict]:
        cp = self.prs.core_properties
        cp.title = self.deck.title
        if self.deck.author:
            cp.author = self.deck.author
        for i, spec in enumerate(self.deck.slides, start=1):
            slide = self.prs.slides.add_slide(self._blank)
            cv = SlideCanvas(self, slide, i)
            getattr(self, f"_{spec.layout}")(cv, spec)
            if spec.notes:
                slide.notes_slide.notes_text_frame.text = spec.notes
        self.prs.save(path)
        return self.issues

    # =================================================================== #
    # Layout archetypes
    # =================================================================== #
    def _title(self, cv: SlideCanvas, s: TitleSlide):
        t, W, H = self.theme, cv.W, cv.H
        cv.bg(t.c("primary"))
        panel_w = 3.1
        cv.rect(W - panel_w, 0, panel_w, H, mix(t.c("primary"), t.c("on_primary"), 0.07))
        cv.rect(W - panel_w, H - 1.6, panel_w, 0.06, t.c("accent"))
        x, w = cv.mx + 0.15, W - panel_w - cv.mx - 0.9
        y = 2.05
        if s.kicker:
            cv.text(x, y, w, 0.4, [TP(s.kicker.upper())], size=cv.sz["kicker"], min_size=10,
                    color=t.c("accent"), bold=True, element="kicker", max_lines=1)
        y += 0.55
        cv.rect(x, y, 0.9, 0.06, t.c("accent"))
        y += 0.35
        f = cv.text(x, y, w, 2.6, [TP(s.title)], size=cv.sz["cover_title"], min_size=30,
                    color=t.c("on_primary"), font=t.heading_font, bold=True, ls=0.95,
                    element="cover title", max_lines=3)
        y += f.height_in + 0.3
        if s.subtitle:
            cv.text(x, y, w, 1.3, [TP(s.subtitle)], size=cv.sz["cover_subtitle"], min_size=14,
                    color=mix(t.c("on_primary"), t.c("primary"), 0.25), element="subtitle",
                    max_lines=3)
        meta = " · ".join(v for v in [self.deck.author, self.deck.date] if v)
        if meta:
            cv.text(x, H - 1.0, w, 0.4, [TP(meta)], size=cv.sz["kicker"], min_size=10,
                    color=mix(t.c("on_primary"), t.c("primary"), 0.35), report=False)

    def _section(self, cv: SlideCanvas, s: SectionSlide):
        t = self.theme
        cv.bg(t.c("surface"))
        cv.rect(0, 0, 0.22, cv.H, t.c("accent"))
        x, w = cv.mx + 0.6, cv.W - 2 * cv.mx - 1.2
        y = 2.3
        if s.number:
            f = cv.text(x, y, w, 1.2, [TP(s.number)], size=72, min_size=40, color=t.c("accent"),
                        font=t.heading_font, bold=True, ls=0.9, max_lines=1, element="section number")
            y += f.height_in + 0.15
        f = cv.text(x, y, w, 2.0, [TP(s.title)], size=cv.sz["section"], min_size=28,
                    color=t.c("ink"), font=t.heading_font, bold=True, ls=0.95,
                    element="section title", max_lines=2)
        y += f.height_in + 0.25
        if s.subtitle:
            cv.text(x, y, w, 1.2, [TP(s.subtitle)], size=cv.sz["sub"] + 2, min_size=13,
                    color=t.c("muted"), element="section subtitle", max_lines=3)

    def _bullets(self, cv: SlideCanvas, s: BulletsSlide):
        t = self.theme
        cv.chrome(s.title)
        bottom = cv.cbottom
        if s.takeaway:
            bottom -= 1.05
            cv.takeaway_bar(s.takeaway, bottom + 0.25)
        cv.text(cv.mx, cv.ctop, cv.cw * 0.86, bottom - cv.ctop, cv.bullet_paras(s.bullets, t.c("ink")),
                size=cv.sz["body"], min_size=cv.sz["body_min"], color=t.c("ink"), element="bullets")

    def _col_paras(self, cv: SlideCanvas, col: Column) -> list[TP]:
        paras = []
        if col.body:
            paras.append(TP(col.body))
        paras += cv.bullet_paras(col.bullets, self.theme.c("ink"))
        if col.body and col.bullets:
            paras[1].space_before = 0.6
        return paras

    def _column(self, cv: SlideCanvas, col: Column, x, y, w, h, head_color: str, card: bool,
                body_pt: float):
        t = self.theme
        pad = 0.35 if card else 0.0
        if card:
            cv.rect(x, y, w, h, t.c("surface"))
            cv.rect(x, y, w, 0.07, head_color)
        hx, hy, hw = x + pad, y + (0.32 if card else 0), w - 2 * pad
        cv.text(hx, hy, hw, 0.5, [TP(col.heading)], size=cv.sz["body"] + 2, min_size=14,
                color=head_color if not card else t.c("ink"), font=t.heading_font, bold=True,
                element="column heading", max_lines=1)
        ry = hy + 0.62
        if not card:
            cv.rect(hx, ry - 0.08, hw, 0.02, t.c("line"))
        ry += 0.12
        cv.text(hx, ry, hw, y + h - ry - pad + 0.05, self._col_paras(cv, col), size=body_pt,
                min_size=body_pt, color=t.c("ink"), element="column body")

    def _two_column(self, cv: SlideCanvas, s: TwoColumnSlide):
        t = self.theme
        cv.chrome(s.title)
        card = s.style != "plain"
        g = t.slide["gutter"] + (0.25 if s.style == "before_after" else 0.1 if card else 0.3)
        w = (cv.cw - g) / 2
        avail = cv.cbottom - cv.ctop
        pad = 0.35 if card else 0.0
        head = (0.32 if card else 0) + 0.74
        # one body size for both columns; card height hugs the taller column
        fits = [textfit.fit(
            [textfit.Para(markup.plain(p.text), p.scale, cv._indent(p), p.bold, p.space_before)
             for p in self._col_paras(cv, c)],
            t.body_font, w - 2 * pad, avail - head - pad, cv.sz["body"] - 2, cv.sz["body_min"] - 1,
            cv.ls) for c in (s.left, s.right)]
        body_pt = min(f.pt for f in fits)
        need = max(f.height_in for f in fits) * body_pt / max(f.pt for f in fits)
        h = min(avail, max(2.6, head + need + pad + 0.35)) if card else avail
        if s.style == "before_after":
            colors = (t.c("muted"), t.c("accent"))
        elif s.style == "compare":
            colors = (t.c("primary"), t.c("primary"))
        else:
            colors = (t.c("primary"), t.c("accent"))
        self._column(cv, s.left, cv.mx, cv.ctop, w, h, colors[0], card, body_pt)
        self._column(cv, s.right, cv.mx + w + g, cv.ctop, w, h, colors[1], card, body_pt)
        if s.style == "before_after":
            ax = cv.mx + w + g / 2 - 0.24
            cv.rect(ax, cv.ctop + h / 2 - 0.24, 0.48, 0.48, t.c("accent"), shape=MSO_SHAPE.RIGHT_ARROW)

    def _big_number(self, cv: SlideCanvas, s: BigNumberSlide):
        t = self.theme
        cv.chrome(s.title)
        h = cv.cbottom - cv.ctop
        if s.context:
            lw = cv.cw * 0.52
        else:
            lw = cv.cw
        y = cv.ctop + 0.35
        f = cv.text(cv.mx, y, lw, 2.0, [TP(s.number)], size=cv.sz["big_number"], min_size=60,
                    color=t.c("accent"), font=t.heading_font, bold=True, ls=0.9, max_lines=1,
                    element="big number")
        y += f.height_in + 0.2
        cv.text(cv.mx, y, lw - 0.3, 1.4, [TP(s.label)], size=cv.sz["body"] + 4, min_size=16,
                color=t.c("ink"), bold=True, element="number label", max_lines=3)
        if s.context:
            rx = cv.mx + lw + 0.5
            rw = cv.cw - lw - 0.5
            cv.rect(rx, cv.ctop + 0.4, 0.05, h - 0.8, t.c("line"))
            cv.text(rx + 0.4, cv.ctop + 0.4, rw - 0.4, h - 0.8, [TP(s.context)],
                    size=cv.sz["body"], min_size=cv.sz["body_min"], color=t.c("muted"),
                    anchor="m", element="context")

    def _kpis(self, cv: SlideCanvas, s: KPISlide):
        t = self.theme
        cv.chrome(s.title)
        n = len(s.items)
        g = 0.3
        w = (cv.cw - g * (n - 1)) / n
        y, h = cv.ctop + 0.15, 2.7
        # shared value size so cards look like a set
        value_pt = min(textfit.fit([textfit.Para(k.value, bold=True)], t.heading_font, w - 0.6, 1.0,
                                   cv.sz["kpi_value"], 24, 1.0, 1).pt for k in s.items)
        for j, k in enumerate(s.items):
            x = cv.mx + j * (w + g)
            cv.rect(x, y, w, h, t.c("surface"))
            cv.rect(x, y, w, 0.07, t.c("accent") if j == 0 else t.c("primary"))
            cv.text(x + 0.3, y + 0.4, w - 0.6, 1.0, [TP(k.value)], size=value_pt, min_size=value_pt,
                    color=t.c("ink"), font=t.heading_font, bold=True, ls=1.0, max_lines=1,
                    element="KPI value")
            cv.text(x + 0.3, y + 1.35, w - 0.6, 0.8, [TP(k.label)], size=cv.sz["kpi_label"] + 2,
                    min_size=11, color=t.c("muted"), element="KPI label", max_lines=2)
            if k.delta:
                sent = k.sentiment or ("negative" if k.delta.strip().startswith(("-", "−", "▼", "↓"))
                                       else "positive")
                col = {"positive": t.c("positive"), "negative": t.c("negative"),
                       "neutral": t.c("muted")}[sent]
                cv.text(x + 0.3, y + h - 0.55, w - 0.6, 0.35, [TP(k.delta)],
                        size=cv.sz["kpi_label"], min_size=10, bold=True, color=col,
                        element="KPI delta", max_lines=1)
        if s.takeaway:
            cv.takeaway_bar(s.takeaway, y + h + 0.45)

    def _chart(self, cv: SlideCanvas, s: ChartSlide):
        t = self.theme
        cv.chrome(s.title)
        h = cv.cbottom - cv.ctop - (0.35 if s.source else 0)
        if s.takeaway:
            side = 3.4
            cw = cv.cw - side - 0.45
            sx = cv.mx + cw + 0.45
            cv.rect(sx, cv.ctop + 0.1, side, 0.06, t.c("accent"))
            cv.text(sx, cv.ctop + 0.4, side, h - 0.5, [TP(s.takeaway)], size=cv.sz["body"] + 2,
                    min_size=14, color=t.c("ink"), bold=False, element="chart takeaway")
        else:
            cw = cv.cw
        add_chart(cv, s.chart, cv.mx, cv.ctop, cw, h)
        if s.source:
            cv.text(cv.mx, cv.cbottom - 0.25, cv.cw, 0.3, [TP(f"Source: {s.source}")],
                    size=cv.sz["small"], min_size=8, color=t.c("muted"), report=False)

    def _table(self, cv: SlideCanvas, s: TableSlide):
        cv.chrome(s.title)
        h = cv.cbottom - cv.ctop - (0.4 if s.source else 0)
        add_table(cv, s.columns, s.rows, cv.mx, cv.ctop, cv.cw, h, s.highlight_row)
        if s.source:
            cv.text(cv.mx, cv.cbottom - 0.25, cv.cw, 0.3, [TP(f"Source: {s.source}")],
                    size=cv.sz["small"], min_size=8, color=self.theme.c("muted"), report=False)

    def _timeline(self, cv: SlideCanvas, s: TimelineSlide):
        t = self.theme
        cv.chrome(s.title)
        n = len(s.steps)
        colw = cv.cw / n
        pad = 0.18
        ly = cv.ctop + 1.0
        cv.rect(cv.mx, ly - 0.015, cv.cw, 0.03, t.c("line"))
        tw = colw - 2 * pad
        title_pt = min(textfit.fit([textfit.Para(st.title, bold=True)], t.body_font, tw, 1.1,
                                   cv.sz["body"], 12, 1.05).pt for st in s.steps)
        det_h = cv.cbottom - (ly + 0.45 + 1.2)
        det_pt = min([textfit.fit([textfit.Para(st.detail)], t.body_font, tw, det_h,
                                  cv.sz["sub"], 10, 1.1).pt for st in s.steps if st.detail] or [12])
        for j, st in enumerate(s.steps):
            x = cv.mx + j * colw
            first = j == 0
            dot = cv.rect(x + pad, ly - 0.14, 0.28, 0.28,
                          t.c("accent") if first else t.c("primary"), shape=MSO_SHAPE.OVAL)
            cv.text(x + pad, ly - 0.75, tw, 0.4, [TP(st.label.upper())], size=cv.sz["kicker"],
                    min_size=10, color=t.c("accent") if first else t.c("muted"), bold=True,
                    max_lines=1, element="timeline label")
            f = cv.text(x + pad, ly + 0.45, tw, 1.15, [TP(st.title)], size=title_pt,
                        min_size=title_pt, color=t.c("ink"), bold=True, ls=1.05,
                        element="timeline title")
            if st.detail:
                cv.text(x + pad, ly + 0.45 + f.height_in + 0.15, tw, det_h, [TP(st.detail)],
                        size=det_pt, min_size=det_pt, color=t.c("muted"), element="timeline detail")

    def _quote(self, cv: SlideCanvas, s: QuoteSlide):
        t = self.theme
        cv.bg(t.c("surface"))
        x, w = cv.mx + 0.9, cv.W - 2 * cv.mx - 1.8
        cv.text(x - 0.1, 0.9, 2, 1.6, [TP("“")], size=150, min_size=150, color=t.c("accent"),
                font=t.heading_font, bold=True, ls=0.8, report=False)
        f = cv.text(x, 2.3, w, 3.3, [TP(s.quote)], size=cv.sz["quote"], min_size=18,
                    color=t.c("ink"), font=t.heading_font, ls=1.15, element="quote")
        if s.attribution:
            cv.rect(x, 2.3 + f.height_in + 0.35, 0.6, 0.04, t.c("accent"))
            cv.text(x, 2.3 + f.height_in + 0.55, w, 0.5, [TP(s.attribution)], size=cv.sz["sub"],
                    min_size=12, color=t.c("muted"), bold=True, max_lines=1, element="attribution")

    def _closing(self, cv: SlideCanvas, s: ClosingSlide):
        t = self.theme
        cv.bg(t.c("primary"))
        x, w = cv.mx + 0.15, cv.W - 2 * cv.mx - 0.3
        y = 1.6 if s.bullets else 2.8
        cv.rect(x, y - 0.35, 0.9, 0.06, t.c("accent"))
        f = cv.text(x, y, w, 1.4, [TP(s.title)], size=cv.sz["cover_title"] - 4, min_size=28,
                    color=t.c("on_primary"), font=t.heading_font, bold=True, ls=0.95,
                    element="closing title", max_lines=2)
        y += f.height_in + 0.45
        if s.bullets:
            paras = [TP(b, bullet="#", space_before=0 if k == 0 else 0.6) for k, b in enumerate(s.bullets)]
            cv.text(x, y, w * 0.8, cv.H - y - 1.2, paras, size=cv.sz["body"], min_size=14,
                    color=t.c("on_primary"), element="next steps")
        if s.contact:
            cv.text(x, cv.H - 0.95, w, 0.4, [TP(s.contact)], size=cv.sz["kicker"] + 1, min_size=10,
                    color=t.c("accent"), bold=True, max_lines=1, element="contact")


# ======================================================================= #
# Charts
# ======================================================================= #
CHART_TYPES = {
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "line": XL_CHART_TYPE.LINE_MARKERS,
    "pie": XL_CHART_TYPE.PIE,
    "doughnut": XL_CHART_TYPE.DOUGHNUT,
    "stacked_column": XL_CHART_TYPE.COLUMN_STACKED,
    "stacked_bar": XL_CHART_TYPE.BAR_STACKED,
}


def add_chart(cv: SlideCanvas, spec: ChartSpec, x, y, w, h):
    t = cv.t
    data = CategoryChartData(number_format=spec.number_format)
    data.categories = spec.categories
    for se in spec.series:
        data.add_series(se.name, se.values)
    gf = cv.s.shapes.add_chart(CHART_TYPES[spec.type], Inches(x), Inches(y), Inches(w), Inches(h), data)
    ch = gf.chart
    ch.has_title = False  # the slide's action title is the chart title
    ch.font.size = Pt(13)
    ch.font.name = t.body_font
    ch.font.color.rgb = _rgb(t.c("muted"))
    pal = [p.lstrip("#") for p in t.chart_palette]
    n_series, n_cat = len(spec.series), len(spec.categories)
    radial = spec.type in ("pie", "doughnut")
    stacked = spec.type.startswith("stacked")

    ch.has_legend = n_series > 1 or radial
    if ch.has_legend:
        ch.legend.position = XL_LEGEND_POSITION.RIGHT if radial else XL_LEGEND_POSITION.TOP
        ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(13)
        ch.legend.font.color.rgb = _rgb(t.c("ink"))

    plot = ch.plots[0]
    if not radial:
        if spec.type in ("column", "bar", "stacked_column", "stacked_bar"):
            plot.gap_width = 55 if n_cat > 4 else 80
            plot.overlap = 100 if stacked else (-10 if n_series > 1 else 0)
        ca, va = ch.category_axis, ch.value_axis
        ca.format.line.color.rgb = _rgb(t.c("line"))
        ca.major_tick_mark = XL_TICK_MARK.NONE
        ca.tick_labels.font.size = Pt(13)
        ca.tick_labels.font.color.rgb = _rgb(t.c("ink"))
        ca.has_major_gridlines = False
        # keep category labels at the edge of the plot, never on the zero line between
        # positive and negative bars
        ca.tick_label_position = XL_TICK_LABEL_POSITION.LOW
        if spec.type in ("bar", "stacked_bar"):
            ca.reverse_order = True   # first category on top, as people read
        clean = spec.data_labels and spec.type != "line" and n_cat <= 14
        va.has_major_gridlines = not clean
        if va.has_major_gridlines:
            va.major_gridlines.format.line.color.rgb = _rgb(t.c("line"))
            va.major_gridlines.format.line.width = Pt(0.75)
        va.visible = not clean
        va.format.line.fill.background()
        va.major_tick_mark = XL_TICK_MARK.NONE
        va.tick_labels.font.size = Pt(12)
        va.tick_labels.number_format = spec.number_format
        va.tick_labels.number_format_is_linked = False
        if spec.y_title and not clean:
            va.has_title = True
            va.axis_title.text_frame.text = spec.y_title
            va.axis_title.text_frame.paragraphs[0].runs[0].font.size = Pt(12)

    for si, series in enumerate(plot.series):
        base = pal[si % len(pal)]
        if spec.type == "line":
            series.smooth = False
            series.format.line.color.rgb = _rgb(base)
            series.format.line.width = Pt(3)
            series.marker.style = 8  # circle
            series.marker.size = 8
            series.marker.format.fill.solid()
            series.marker.format.fill.fore_color.rgb = _rgb(base)
            series.marker.format.line.color.rgb = _rgb(t.c("bg"))
            continue
        if radial:
            for pi in range(n_cat):
                pt = series.points[pi]
                pt.format.fill.solid()
                col = pal[pi % len(pal)]
                if spec.highlight is not None and pi != spec.highlight:
                    col = mix(col, t.c("bg"), 0.55)
                pt.format.fill.fore_color.rgb = _rgb(col)
                pt.format.line.color.rgb = _rgb(t.c("bg"))
                pt.format.line.width = Pt(1.5)
            continue
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = _rgb(base)
        if spec.highlight is not None:
            for pi in range(n_cat):
                pt = series.points[pi]
                pt.format.fill.solid()
                if n_series == 1:
                    col = t.c("accent") if pi == spec.highlight else t.c("chart_muted")
                else:
                    col = base if pi == spec.highlight else mix(base, t.c("bg"), 0.6)
                pt.format.fill.fore_color.rgb = _rgb(col)
        elif n_series == 1:
            series.format.fill.fore_color.rgb = _rgb(pal[0])

    if spec.data_labels and not (spec.type == "line" and n_cat > 8):
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.number_format = spec.number_format
        dl.number_format_is_linked = False
        dl.font.size = Pt(13 if n_cat <= 8 else 11)
        dl.font.bold = True
        if stacked:
            dl.position = XL_LABEL_POSITION.CENTER
            dl.font.color.rgb = _rgb("FFFFFF")
        elif spec.type in ("column", "bar"):
            dl.position = XL_LABEL_POSITION.OUTSIDE_END
            dl.font.color.rgb = _rgb(t.c("ink"))
        elif spec.type == "line":
            dl.position = XL_LABEL_POSITION.ABOVE
            dl.font.color.rgb = _rgb(t.c("ink"))
        elif spec.type == "pie":
            dl.position = XL_LABEL_POSITION.OUTSIDE_END
            dl.font.color.rgb = _rgb(t.c("ink"))
        else:
            dl.font.color.rgb = _rgb("FFFFFF")
    if spec.type == "doughnut":
        hole = plot._element.find(qn("c:holeSize"))
        if hole is None:
            hole = etree.SubElement(plot._element, qn("c:holeSize"))
        hole.set("val", "58")
    return ch


# ======================================================================= #
# Tables
# ======================================================================= #
NO_GRID_STYLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"


def _cell_border_bottom(cell, color: str, width_pt: float):
    tcPr = cell._tc.get_or_add_tcPr()
    ln = etree.Element(qn("a:lnB"), w=str(int(Pt(width_pt))), cap="flat", cmpd="sng", algn="ctr")
    sf = etree.SubElement(ln, qn("a:solidFill"))
    etree.SubElement(sf, qn("a:srgbClr"), val=color)
    etree.SubElement(ln, qn("a:prstDash"), val="solid")
    tcPr.insert(0, ln)


def add_table(cv: SlideCanvas, columns, rows, x, y, w, h, highlight_row=None):
    t = cv.t
    font = t.body_font
    ncol, nrow = len(columns), len(rows) + 1
    # column widths ~ content length (dampened), with a floor
    weights = []
    for c in range(ncol):
        longest = max([len(columns[c])] + [len(r[c]) for r in rows])
        weights.append(max(6, min(longest, 48)) ** 0.75)
    widths = [w * wt / sum(weights) for wt in weights]
    numeric = [sum(bool(NUMERIC.match(r[c])) for r in rows) >= max(1, len(rows) * 0.6) for c in range(ncol)]

    pad_x, pad_y = 0.12, 0.07

    def row_heights(pt):
        hs = []
        for ri, cells in enumerate([columns] + rows):
            best = 0
            for c, txt in enumerate(cells):
                lines = textfit.wrap(markup.plain(txt), font, pt, widths[c] - 2 * pad_x, bold=ri == 0)
                best = max(best, len(lines))
            hs.append(best * pt * textfit.LINE_HEIGHT / 72 + 2 * pad_y + 0.06)
        return hs

    pt = t.slide["size"]["table"]
    while True:
        hs = row_heights(pt)
        if sum(hs) <= h or pt <= 10:
            break
        pt -= 1
    if sum(hs) > h:
        cv.r.issue(cv.i, "error", "overflow", f"table is too tall even at {pt}pt.",
                   "Cut rows (≤7 is ideal), shorten cell text, or split into two slides.")

    gf = cv.s.shapes.add_table(nrow, ncol, Inches(x), Inches(y), Inches(w), Inches(sum(hs)))
    tbl = gf.table
    sid = tbl._tbl.tblPr.find(qn("a:tableStyleId"))
    if sid is None:
        sid = etree.SubElement(tbl._tbl.tblPr, qn("a:tableStyleId"))
    sid.text = NO_GRID_STYLE
    tbl.first_row = True
    tbl.horz_banding = False
    for c, cw in enumerate(widths):
        tbl.columns[c].width = Inches(cw)
    for ri, rh in enumerate(hs):
        tbl.rows[ri].height = Inches(rh)

    for ri, cells in enumerate([columns] + rows):
        header = ri == 0
        hl = highlight_row is not None and ri - 1 == highlight_row
        for c, txt in enumerate(cells):
            cell = tbl.cell(ri, c)
            _cell_border_bottom(cell, t.c("primary") if header else t.c("line"), 1.25 if header else 0.75)
            cell.fill.solid()
            if header:
                cell.fill.fore_color.rgb = _rgb(t.c("bg"))
            elif hl:
                cell.fill.fore_color.rgb = _rgb(mix(t.c("accent"), t.c("bg"), 0.86))
            else:
                cell.fill.fore_color.rgb = _rgb(t.c("bg"))
            cell.margin_left = cell.margin_right = Inches(pad_x)
            cell.margin_top = cell.margin_bottom = Inches(pad_y)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.RIGHT if numeric[c] and c > 0 else PP_ALIGN.LEFT
            for seg, b, it in markup.runs(txt):
                run = p.add_run()
                run.text = seg
                run.font.name = font
                run.font.size = Pt(pt if not header else max(pt - 1, 10))
                run.font.bold = header or hl or b or (c == 0 and not header)
                run.font.italic = it
                run.font.color.rgb = _rgb(t.c("muted") if header else t.c("ink"))
    return tbl


def render_deck(deck: Deck, theme: Theme, path: str) -> list[dict]:
    return DeckRenderer(deck, theme).render(path)
