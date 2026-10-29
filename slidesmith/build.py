"""validate -> render -> rasterize -> lint, in one call.

This is the single tool the orchestrating agent needs. It always writes a
report.json/report.md, even when validation fails, so the agent has something
concrete to fix.
"""
from __future__ import annotations

import json
import tempfile
import traceback
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Optional

from pydantic import ValidationError

from . import lint, provenance
from .rasterize import RasterizeError, contact_sheet, render_pages
from .render_docx import render_doc
from .render_pptx import render_deck
from .schema import Deck, load_spec
from .theme import load_theme


@dataclass
class BuildResult:
    ok: bool
    kind: Optional[str] = None
    output: Optional[str] = None
    pdf: Optional[str] = None
    pages: list[str] = field(default_factory=list)
    contact_sheet: Optional[str] = None
    renderer: Optional[str] = None
    report: Optional[str] = None
    summary: dict = field(default_factory=dict)

    def to_dict(self):
        return asdict(self)


def format_validation_error(e: ValidationError) -> list[dict]:
    """Turn pydantic paths like ('slides', 2, 'chart', 'chart') into 'slide 3 › chart'."""
    out = []
    for err in e.errors():
        parts = list(err["loc"])
        where, index = "spec", 0
        if len(parts) >= 2 and parts[0] in ("slides", "blocks") and isinstance(parts[1], int):
            index = parts[1] + 1
            where = f"{'slide' if parts[0] == 'slides' else 'block'} {index}"
            parts = parts[2:]
            if parts and isinstance(parts[0], str):  # discriminator tag, e.g. 'chart'
                parts = parts[1:]
        field = ".".join(str(x) for x in parts)
        if field:
            where += f" › {field}"
        out.append({"where": where, "index": index, "severity": "error", "rule": "schema",
                    "message": err["msg"], "hint": f"Fix the spec at {where}."})
    return out


def writable_path(p: Path) -> Path:
    """p, or 'name (2).pptx' etc. if p is locked — PowerPoint/Word lock files they have open."""
    if not p.exists():
        return p
    try:
        with open(p, "r+b"):
            return p
    except PermissionError:
        for n in range(2, 100):
            alt = p.with_name(f"{p.stem} ({n}){p.suffix}")
            if not alt.exists():
                return alt
            try:
                with open(alt, "r+b"):
                    return alt
            except PermissionError:
                continue
        raise


def read_spec(path: str | Path) -> dict:
    return json.loads(Path(path).read_text(encoding="utf-8-sig"))


def build(spec_path: str | Path, out_dir: str | Path | None = None, theme: str | None = None,
          rasterize: bool = True, dpi: int = 90, renderer: str = "auto") -> BuildResult:
    """Never raises: any failure comes back as a report the agent can act on."""
    try:
        return _build(spec_path, out_dir, theme, rasterize, dpi, renderer)
    except Exception as e:  # noqa: BLE001
        tb = traceback.format_exc(limit=6)
        try:
            sp = Path(spec_path)
            if out_dir:
                out = Path(out_dir)
            elif sp.exists():
                out = sp.parent / "build" / sp.name.split(".")[0]
            else:
                raise FileNotFoundError
            out.mkdir(parents=True, exist_ok=True)
        except Exception:  # noqa: BLE001
            out = Path(tempfile.mkdtemp(prefix="slidesmith-"))
        return _fail(out, [{"where": "global", "index": 0, "severity": "error", "rule": "internal",
                            "message": f"Build failed unexpectedly: {type(e).__name__}: {e}",
                            "hint": "Try simplifying the slide/block named in the trace, or theme "
                                    "'default'. Trace:\n" + tb}])


def _build(spec_path, out_dir, theme, rasterize, dpi, renderer) -> BuildResult:
    spec_path = Path(spec_path)
    if not spec_path.exists():
        raise FileNotFoundError(f"spec not found: {spec_path}")
    stem = spec_path.name.split(".")[0]
    out = Path(out_dir) if out_dir else spec_path.parent / "build" / stem
    out.mkdir(parents=True, exist_ok=True)

    try:
        data = read_spec(spec_path)
        spec = load_spec(data)
    except json.JSONDecodeError as e:
        return _fail(out, [{"where": f"line {e.lineno}", "index": 0, "severity": "error",
                            "rule": "json", "message": str(e), "hint": "Fix the JSON syntax."}])
    except ValidationError as e:
        return _fail(out, format_validation_error(e))

    try:
        th = load_theme(theme or spec.theme)
    except FileNotFoundError as e:
        return _fail(out, [{"where": "spec › theme", "index": 0, "severity": "error", "rule": "theme",
                            "message": str(e), "hint": "Use a listed theme name or a path to a theme .json."}])
    kind = "deck" if isinstance(spec, Deck) else "doc"
    wanted = out / f"{stem}.{'pptx' if kind == 'deck' else 'docx'}"
    target = writable_path(wanted)
    pre: list[dict] = []
    if target != wanted:
        pre.append({"where": "global", "index": 0, "severity": "warning", "rule": "file_locked",
                    "message": f"{wanted.name} is open in another program (Windows locks it), "
                               f"so this build was written to {target.name}.",
                    "hint": "Close the file in PowerPoint/Word to overwrite it on the next build."})

    if kind == "deck":
        issues = pre + render_deck(spec, th, str(target)) + lint.lint_deck(spec, th)
    else:
        issues = pre + render_doc(spec, th, str(target)) + lint.lint_doc(spec, th)

    try:
        issues += provenance.check(spec, spec_path)
    except Exception as e:  # noqa: BLE001
        issues.append({"where": "global", "index": 0, "severity": "warning", "rule": "provenance",
                       "message": f"figure check failed: {e}", "hint": "Verify numbers manually."})

    # the file must open again cleanly before we call it done
    try:
        if kind == "deck":
            from pptx import Presentation
            n = len(Presentation(str(target)).slides)
            if n != len(spec.slides):
                raise ValueError(f"expected {len(spec.slides)} slides, file has {n}")
        else:
            from docx import Document as _Doc
            _Doc(str(target))
    except Exception as e:  # noqa: BLE001
        return _fail(out, [{"where": "global", "index": 0, "severity": "error", "rule": "corrupt_output",
                            "message": f"The generated file did not re-open: {e}",
                            "hint": "Do not use this file. Rebuild; if it persists, report the spec."}])

    res = BuildResult(ok=True, kind=kind, output=str(target))
    if rasterize:
        try:
            rr = render_pages(target, out / "pages", dpi=dpi, renderer=renderer, spec=spec, theme=th)
            pages = rr.pages
            res.renderer = rr.renderer
            res.pdf = str(rr.pdf) if rr.pdf else None
            res.pages = [str(p) for p in pages]
            for note in rr.notes:
                issues.append({"where": "global", "index": 0, "severity": "info", "rule": "renderer",
                               "message": note, "hint": ""})
            res.contact_sheet = str(contact_sheet(pages, out / "contact_sheet.png",
                                                  cols=3 if kind == "deck" else 4,
                                                  thumb_w=560 if kind == "deck" else 420))
            issues += lint.lint_pages(kind, pages, th.colors["bg"])
        except Exception as e:  # noqa: BLE001 — never lose the built file over a preview problem
            issues.append({"where": "global", "index": 0, "severity": "warning", "rule": "raster",
                           "message": str(e), "hint": "Visual review is unavailable for this build."})
    res.summary = lint.summarize(issues)
    res.ok = res.summary["errors"] == 0
    res.report = _write_report(out, res)
    return res


def _fail(out: Path, issues: list[dict]) -> BuildResult:
    res = BuildResult(ok=False, summary=lint.summarize(issues))
    res.report = _write_report(out, res)
    return res


def _write_report(out: Path, res: BuildResult) -> str:
    (out / "report.json").write_text(json.dumps(res.to_dict(), indent=2, ensure_ascii=False), encoding="utf-8")
    s = res.summary
    lines = [f"# Build report — score {s['score']}/100",
             f"{s['errors']} errors · {s['warnings']} warnings · {s['info']} info", ""]
    if res.output:
        lines += [f"- Output: `{res.output}`"]
    if res.contact_sheet:
        lines += [f"- Contact sheet: `{res.contact_sheet}`",
                  f"- Pages: {len(res.pages)} PNGs in `pages/` (rendered by: {res.renderer})"]
    lines.append("")
    for sev in ("error", "warning", "info"):
        group = [x for x in s["issues"] if x["severity"] == sev]
        if group:
            lines.append(f"## {sev.capitalize()}s")
            for x in group:
                hint = f" → {x['hint']}" if x.get("hint") else ""
                lines.append(f"- **{x['where']}** [{x['rule']}] {x['message']}{hint}")
            lines.append("")
    p = out / "report.md"
    p.write_text("\n".join(lines), encoding="utf-8")
    return str(p)
