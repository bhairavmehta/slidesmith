"""Number provenance: is every figure in the deck/doc backed by a source file?

The most damaging failure of an AI-written deck is a plausible but wrong number. When a
spec lists its `sources` (paths relative to the spec), every number that appears in the
spec's text and chart data is looked up in those files. Numbers that cannot be found are
reported as `unverified_figure` warnings so a human checks them before anything is sent.

Matching is tolerant of formatting: 86% ↔ 0.86 ↔ 86, $4.2M ↔ 4,200,000 ↔ 4.2, and values
rounded in the spec (0.82 vs 0.8213 in the source). Derived numbers (differences, ratios)
will not be found — that is intended: they get flagged so someone checks the arithmetic.

Readable sources: .txt .md .csv .tsv .json .docx .pptx .xlsx/.xlsm (openpyxl) .pdf (pypdfium2).
"""
from __future__ import annotations

import json
import math
import re
from pathlib import Path

from pydantic import BaseModel

NUM = re.compile(
    r"(?<![\w.])([-+−]?)\s?([$€£¥])?\s?(\d{1,3}(?:,\d{3})+|\d+)(\.\d+)?\s?(%|pp|pts?|bps|[kKmMbB]n?|x|×)?(?![\w])"
)
SCALE = {"k": 1e3, "m": 1e6, "mn": 1e6, "b": 1e9, "bn": 1e9}


def _tokens(text: str) -> list[tuple[float, int, str]]:
    """(value, decimals, unit) for every number in text."""
    out = []
    for m in NUM.finditer(text or ""):
        sign, cur, whole, frac, unit = m.groups()
        try:
            v = float(whole.replace(",", "") + (frac or ""))
        except ValueError:
            continue
        if sign in ("-", "−"):
            v = -v
        out.append((v, len(frac) - 1 if frac else 0, (unit or "").lower()))
    return out


def _variants(v: float, unit: str) -> set[float]:
    """Equivalent readings of one source number."""
    vs = {v, abs(v)}
    if unit == "%":
        vs |= {v / 100, abs(v) / 100}
    s = SCALE.get(unit)
    if s:
        vs |= {v * s, abs(v * s)}
    # a plain source number may be shown as % or in K/M/B in the spec
    vs |= {v * 100, abs(v) * 100}
    for s in (1e3, 1e6, 1e9):
        vs |= {v / s, abs(v) / s}
    return vs


# --------------------------------------------------------------------------- #
def read_source(p: Path) -> str:
    suf = p.suffix.lower()
    if suf in (".txt", ".md", ".csv", ".tsv", ".json", ".yaml", ".yml"):
        return p.read_text(encoding="utf-8-sig", errors="replace")
    if suf == ".docx":
        from docx import Document

        d = Document(str(p))
        parts = [para.text for para in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts += [c.text for c in row.cells]
        return "\n".join(parts)
    if suf == ".pptx":
        from pptx import Presentation

        parts = []
        for slide in Presentation(str(p)).slides:
            for shp in slide.shapes:
                if shp.has_text_frame:
                    parts.append(shp.text_frame.text)
                if getattr(shp, "has_table", False) and shp.has_table:
                    parts += [c.text for r in shp.table.rows for c in r.cells]
                if getattr(shp, "has_chart", False) and shp.has_chart:
                    for plot in shp.chart.plots:
                        for s in plot.series:
                            parts += [repr(v) for v in s.values if v is not None]
        return "\n".join(parts)
    if suf in (".xlsx", ".xlsm"):
        import openpyxl

        wb = openpyxl.load_workbook(str(p), read_only=True, data_only=True)
        parts = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                parts += [repr(c) if isinstance(c, float) else str(c) for c in row if c is not None]
        return "\n".join(parts)
    if suf == ".pdf":
        import pypdfium2 as pdfium

        doc = pdfium.PdfDocument(str(p))
        try:
            return "\n".join(doc[i].get_textpage().get_text_range() for i in range(len(doc)))
        finally:
            doc.close()
    raise ValueError(f"unsupported source type {suf}")


def _spec_strings(obj, path="") -> list[tuple[str, str]]:
    """(location, text) for every string in the spec, plus chart values."""
    out = []
    if isinstance(obj, BaseModel):
        obj = obj.model_dump()
    if isinstance(obj, dict):
        if "categories" in obj and "series" in obj:  # chart data: check the values themselves
            for s in obj["series"]:
                for i, v in enumerate(s["values"]):
                    cat = obj["categories"][i] if i < len(obj["categories"]) else i
                    out.append((f"{path}chart '{s['name']}' / {cat}", repr(float(v))))
            obj = {k: v for k, v in obj.items() if k not in ("series", "number_format")}
        for k, v in obj.items():
            if k in ("layout", "type", "kind", "theme", "style", "number_format", "sources", "label") and \
                    not (k == "label" and isinstance(v, str) and len(v) > 12):
                continue
            out += _spec_strings(v, f"{path}{k}.")
    elif isinstance(obj, list):
        for i, v in enumerate(obj):
            out += _spec_strings(v, f"{path}{i}.")
    elif isinstance(obj, str):
        out.append((path.rstrip("."), obj))
    return out


def _ignorable(v: float, dec: int, unit: str, text: str) -> bool:
    if dec == 0 and not unit and 0 <= v <= 10:
        return True                      # small counts: "3 steps", "two of 3"
    if dec == 0 and not unit and 1900 <= v <= 2100:
        return True                      # years
    return False


def _matches(v: float, dec: int, pool: list[float]) -> bool:
    for s in pool:
        if math.isclose(v, s, rel_tol=1e-9, abs_tol=1e-12):
            return True
        # spec shows a rounded version of the source value
        if dec >= 0 and abs(round(s, dec) - v) < 10 ** (-dec) * 0.5001:
            if s != 0 and abs(v - s) / max(abs(s), 1e-12) < 0.05:
                return True
    return False


def check(spec, spec_path: Path) -> list[dict]:
    sources = getattr(spec, "sources", None) or []
    kind = "slide" if hasattr(spec, "slides") else "block"
    issues = []

    strings = _spec_strings(spec.slides if kind == "slide" else spec.blocks, "")
    numbers = []
    for loc, text in strings:
        for v, dec, unit in _tokens(text):
            if not _ignorable(v, dec, unit, text):
                numbers.append((loc, text, v, dec, unit))

    if not sources:
        if numbers:
            issues.append({"where": "global", "index": 0, "severity": "info", "rule": "no_sources",
                           "message": f"{len(numbers)} figures in this {'deck' if kind == 'slide' else 'document'}"
                                      " cannot be verified because the spec lists no `sources`.",
                           "hint": "Add \"sources\": [\"data.xlsx\", \"notes.md\"] (paths relative to the spec)."})
        return issues

    pool: list[float] = []
    for src in sources:
        p = (spec_path.parent / src) if not Path(src).is_absolute() else Path(src)
        try:
            text = read_source(p)
        except FileNotFoundError:
            issues.append({"where": "global", "index": 0, "severity": "error", "rule": "source_missing",
                           "message": f"source not found: {src}", "hint": "Fix the path in `sources`."})
            continue
        except Exception as e:  # noqa: BLE001
            issues.append({"where": "global", "index": 0, "severity": "warning", "rule": "source_unreadable",
                           "message": f"could not read source {src}: {e}",
                           "hint": "Export it to .xlsx/.csv/.md/.pdf, or remove it from `sources`."})
            continue
        for v, _dec, unit in _tokens(text):
            pool += list(_variants(v, unit))

    seen = set()
    for loc, text, v, dec, unit in numbers:
        m = re.match(r"(\d+)\.", loc)
        idx = int(m.group(1)) + 1 if m else 0
        variants = [v] + ([v / 100] if unit == "%" else []) + \
                   ([v * SCALE[unit]] if unit in SCALE else [])
        if any(_matches(x, dec, pool) for x in variants):
            continue
        key = (idx, round(v, 6), unit)
        if key in seen:
            continue
        seen.add(key)
        shown = text if len(text) <= 70 else text[:67] + "…"
        issues.append({"where": f"{kind} {idx}" if idx else "global", "index": idx, "severity": "warning",
                       "rule": "unverified_figure",
                       "message": f"figure {v:g}{unit} not found in sources ({loc}: \"{shown}\")",
                       "hint": "Check it against the source. If it is derived (a difference, ratio, "
                               "rounding), show the calculation in the speaker notes; otherwise fix it."})
    return issues
